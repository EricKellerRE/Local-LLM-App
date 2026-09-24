from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable


MARKDOWN_URL_PATTERN = re.compile(r"\[[^\]]*\]\((https?://[^\s]+)\)", re.IGNORECASE)
BARE_URL_PATTERN = re.compile(r"https?://[^\s<>\]]+", re.IGNORECASE)
INLINE_PATTERN = re.compile(
    r"(\[[^\]]+\]\(https?://[^\s]+\)|\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))"
)


def report_word_count(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text))


def unique_urls(text: str) -> list[str]:
    candidates = [*MARKDOWN_URL_PATTERN.findall(text), *BARE_URL_PATTERN.findall(text)]
    urls: list[str] = []
    for candidate in candidates:
        url = candidate.rstrip(".,;:!?")
        while url.endswith(")") and url.count("(") < url.count(")"):
            url = url[:-1]
        if url not in urls:
            urls.append(url)
    return urls


def assemble_research_report(title: str, sections: Iterable[tuple[str, str]]) -> str:
    parts = [f"# {title.strip()}"]
    for section_title, content in sections:
        cleaned = content.strip()
        cleaned = re.sub(r"\A#\s+[^\n]+\n+", "", cleaned)
        if not re.match(r"^#{2,6}\s", cleaned):
            parts.append(f"## {section_title.strip()}")
        parts.append(cleaned)
    body = "\n\n".join(part for part in parts if part.strip()).strip()
    urls = unique_urls(body)
    if urls:
        references = ["## Source URLs", *[f"{index}. [{url}]({url})" for index, url in enumerate(urls, 1)]]
        body = f"{body}\n\n" + "\n".join(references)
    return body + "\n"


def _set_cell_borders(cell, color: str = "D9D9D9") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def _shade_cell(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _configure_table_row(row, *, repeat_header: bool = False) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    properties.append(cannot_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        properties.append(header)


def _add_hyperlink(paragraph, label: str, url: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_inline(paragraph, text: str) -> None:
    text = re.sub(r"\$([^$]+)\$", lambda match: _plain_math(match.group(1)), text)
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^\s]+)\)", token)
        if link:
            _add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _plain_math(text: str) -> str:
    superscripts = str.maketrans("0123456789+-()", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁽⁾")
    subscripts = str.maketrans("0123456789+-()", "₀₁₂₃₄₅₆₇₈₉₊₋₍₎")
    text = re.sub(r"\^\{([^}]+)\}", lambda match: match.group(1).translate(superscripts), text)
    text = re.sub(r"_\{([^}]+)\}", lambda match: match.group(1).translate(subscripts), text)
    return text.replace("\\mathrm", "").replace("{", "").replace("}", "")


def _is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _table_rows(lines: list[str], start: int) -> tuple[list[list[str]], int] | None:
    if start + 1 >= len(lines) or "|" not in lines[start] or not _is_table_separator(lines[start + 1]):
        return None
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and "|" in lines[index] and lines[index].strip():
        if index != start + 1:
            rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
        index += 1
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows], index


def build_report_docx(markdown: str, output_path: Path, *, title: str | None = None) -> Path:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (
        ("Title", 24, 0, 14),
        ("Heading 1", 17, 14, 7),
        ("Heading 2", 14, 11, 5),
        ("Heading 3", 12, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        paragraph_properties = style._element.get_or_add_pPr()
        border = paragraph_properties.find(qn("w:pBdr"))
        if border is not None:
            paragraph_properties.remove(border)

    lines = markdown.replace("\r\n", "\n").split("\n")
    inferred_title = title
    if not inferred_title:
        inferred_title = next((line[2:].strip() for line in lines if line.startswith("# ")), "Research Report")
    title_paragraph = document.add_paragraph(style="Title")
    title_properties = title_paragraph._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_inline(title_paragraph, inferred_title)
    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(16)
    meta_run = metadata.add_run(
        f"Generated locally {datetime.now(timezone.utc).strftime('%B %d, %Y')}  |  "
        f"{report_word_count(markdown):,} words  |  {len(unique_urls(markdown))} source URLs"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(89, 89, 89)

    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        paragraph = document.add_paragraph()
        _add_inline(paragraph, " ".join(part.strip() for part in paragraph_buffer))
        paragraph_buffer.clear()

    index = 0
    skipped_first_title = False
    while index < len(lines):
        line = lines[index].rstrip()
        table_data = _table_rows(lines, index)
        if table_data:
            flush_paragraph()
            rows, index = table_data
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.autofit = True
            for row_index, values in enumerate(rows):
                _configure_table_row(table.rows[row_index], repeat_header=row_index == 0)
                for column_index, value in enumerate(values):
                    cell = table.cell(row_index, column_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    _set_cell_borders(cell)
                    if row_index == 0:
                        _shade_cell(cell, "1F4E78")
                    elif row_index % 2 == 0:
                        _shade_cell(cell, "EEF4F8")
                    paragraph = cell.paragraphs[0]
                    if row_index == 0:
                        paragraph.paragraph_format.keep_with_next = True
                    _add_inline(paragraph, value)
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            if level == 1 and not skipped_first_title:
                skipped_first_title = True
            else:
                paragraph = document.add_paragraph(style=f"Heading {min(max(level - 1, 1), 3)}")
                _add_inline(paragraph, heading.group(2).strip())
            index += 1
            continue
        if re.fullmatch(r"\s*([-*_])(?:\s*\1){2,}\s*", line):
            flush_paragraph()
            index += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*(\d+)[.)]\s+(.+)$", line)
        if bullet or numbered:
            flush_paragraph()
            if bullet:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.28)
                paragraph.paragraph_format.first_line_indent = Inches(-0.22)
                paragraph.add_run("•\u00a0")
                _add_inline(paragraph, bullet.group(1))
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.28)
                paragraph.paragraph_format.first_line_indent = Inches(-0.22)
                paragraph.add_run(f"{numbered.group(1)}.\u00a0").bold = True
                _add_inline(paragraph, numbered.group(2))
            index += 1
            continue
        if not line.strip():
            flush_paragraph()
        else:
            paragraph_buffer.append(line)
        index += 1
    flush_paragraph()

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.widow_control = True
    document.save(output_path)
    return output_path
